#!/usr/bin/env python3
"""
Build clean single-column probate templates for Ginsberg Shulman.

Mirrors the guardianship rebuild (build_guardianship_templates.py):
  - Single-column caption via 2-column borderless table
  - Real Word numbering (numPr/numId=1) per the docx-numbering skill —
    no hardcoded "1.\\t" text runs, 1.5 line spacing
  - docxtemplater conditional blocks instead of "strike each statement"
  - Real Word tables for beneficiaries
  - Broward AI certification above the signature block (renders only for
    Broward-county matters)

Shared helpers (_add_para, _pleading_para, _inject_numbering_part, etc.)
are imported from build_guardianship_templates to avoid duplication.

Currently builds:
  P3-0100  Petition for Administration (testate, FL resident, single petitioner)
"""

import os

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_guardianship_templates import (
    FONT, BODY_SIZE, SMALL_SIZE, TEMPLATE_DIR,
    _add_para, _add_run, _set_run,
    _borderless_table, _table_with_borders, _clear_cell, _cell_para,
    _apply_page_setup, _apply_running_header,
    _ensure_pleading_numbering, _pleading_para,
    _inject_numbering_part,
    _add_broward_ai_certification,
)


# ---------------------------------------------------------------------------
# Probate-specific helpers
# ---------------------------------------------------------------------------

def _add_probate_caption(doc, *, decedent_tag='{decedent_full_name}'):
    """
    Probate caption:
      IN THE CIRCUIT COURT FOR {county} COUNTY, FLORIDA (centered)
      4-row borderless table:
        IN RE: ESTATE OF          |  PROBATE DIVISION
        {decedent_tag},           |
        Deceased.                 |  File No. {file_no}
                                  |  Division {division}
    """
    _add_para(doc, 'IN THE CIRCUIT COURT FOR {county} COUNTY, FLORIDA',
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    caption = _borderless_table(doc, rows=4, cols=2, col_widths_in=[3.5, 3.0])
    _clear_cell(caption.cell(0, 0)); _cell_para(caption.cell(0, 0), 'IN RE: ESTATE OF', space_after=0)
    _clear_cell(caption.cell(0, 1)); _cell_para(caption.cell(0, 1), 'PROBATE DIVISION', space_after=0)
    _clear_cell(caption.cell(1, 0)); _cell_para(caption.cell(1, 0), f'{decedent_tag},', space_after=0)
    _clear_cell(caption.cell(1, 1)); _cell_para(caption.cell(1, 1), '', space_after=0)
    _clear_cell(caption.cell(2, 0)); _cell_para(caption.cell(2, 0), 'Deceased.', space_after=0)
    _clear_cell(caption.cell(2, 1)); _cell_para(caption.cell(2, 1), 'File No. {file_no}', space_after=0)
    _clear_cell(caption.cell(3, 0)); _cell_para(caption.cell(3, 0), '', space_after=0)
    _clear_cell(caption.cell(3, 1)); _cell_para(caption.cell(3, 1), 'Division {division}', space_after=0)
    _add_para(doc, '', space_after=18)


def _beneficiaries_table(doc):
    """Four-column bordered table for probate beneficiaries with a
    docxtemplater paragraph-loop row."""
    tbl = _table_with_borders(doc, rows=2, cols=4, col_widths_in=[1.8, 2.4, 1.4, 0.9])
    for cell in tbl.rows[0].cells:
        _clear_cell(cell)
    _cell_para(tbl.cell(0, 0), 'NAME', bold=True, space_after=0)
    _cell_para(tbl.cell(0, 1), 'ADDRESS', bold=True, space_after=0)
    _cell_para(tbl.cell(0, 2), 'RELATIONSHIP', bold=True, space_after=0)
    _cell_para(tbl.cell(0, 3), 'BIRTH YR.', bold=True, space_after=0)
    for cell in tbl.rows[1].cells:
        _clear_cell(cell)
    _cell_para(tbl.cell(1, 0), '{#beneficiaries}{ben_name}', space_after=0)
    _cell_para(tbl.cell(1, 1), '{ben_address}', space_after=0)
    _cell_para(tbl.cell(1, 2), '{ben_relationship}', space_after=0)
    _cell_para(tbl.cell(1, 3), '{ben_year_of_birth}{/beneficiaries}', space_after=0)
    return tbl


def _add_probate_signature_block(doc):
    """Probate petition signature block: petitioner sig line + attorney sig
    block. Differs from guardianship's (which uses 'Petitioner' role label)
    only in that probate petitions are typically signed by the proposed PR,
    but since petitioner and PR are often the same person we use the generic
    petitioner block.
    """
    indent = Inches(0.5)
    _add_para(doc, 'Signed on this _____ day of {signing_month} {signing_year}.',
              first_indent=indent, space_after=24)
    _add_para(doc, '_______________________________________', space_after=0)
    _add_para(doc, '{petitioner_name}, Petitioner', space_after=24)
    _add_para(doc, '_______________________________________', space_after=0)
    _add_para(doc, '{attorney_name}, Attorney for Petitioner', space_after=18)
    _add_para(doc, 'Email Addresses:', space_after=0)
    _add_para(doc, '{attorney_email}', space_after=0)
    _add_para(doc, '{#attorney_email_secondary}{attorney_email_secondary}{/attorney_email_secondary}', space_after=0)
    _add_para(doc, 'Florida Bar No. {attorney_bar_no}', space_after=12)
    _add_para(doc, '{attorney_firm}', space_after=0)
    _add_para(doc, '{attorney_address}', space_after=12)
    _add_para(doc, 'Telephone {attorney_phone}', space_after=0)


# ---------------------------------------------------------------------------
# P3-0100 Petition for Administration (testate, FL resident, single petitioner)
# ---------------------------------------------------------------------------

def build_p3_0100():
    """Rebuilt per docx-numbering skill + firm conventions.

    Replaces the FLSSI 'strike each statement that is not applicable'
    language with clean docxtemplater conditionals:
      {#petitioner_has_prior_conviction}   item 6 true branch
      {#higher_preference_exists}          item 7 true branch
      {#estate_tax_return_required}        item 9 true branch
      {#domiciliary_proceedings_pending}   item 10 true branch
      will_status ('original' / 'authenticated_other' / 'authenticated_notarial')
                                           item 11 three-way branch
    """
    doc = Document()
    _apply_page_setup(doc)
    _apply_running_header(doc, 'Estate of {decedent_name}')
    _ensure_pleading_numbering(doc)

    _add_probate_caption(doc)

    # ---- Title ----
    _add_para(doc, 'PETITION FOR ADMINISTRATION',
              align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=0)
    _add_para(doc, '(Testate Florida Resident — Single Petitioner)',
              align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, space_after=18)

    # Intro — unnumbered.
    _add_para(doc, 'Petitioner, {petitioner_name}, alleges:', space_after=12)

    # 1. Interest + address
    _pleading_para(doc,
        'Petitioner has an interest in the above estate as {petitioner_interest}. '
        'Petitioner\u2019s address is {petitioner_address}, and the name and office address of petitioner\u2019s attorney are set forth at the end of this petition.')

    # 2. Decedent
    _pleading_para(doc,
        'Decedent, {decedent_full_name}, whose last known address was {decedent_address}, '
        'and the last four digits of whose social security number are {decedent_ssn_last4}, '
        'died on {decedent_death_date} at {decedent_death_place}. On the date of death, '
        'decedent was domiciled in {decedent_domicile} County, Florida.')

    # 3. Beneficiaries
    _pleading_para(doc,
        'So far as is known, the names of the beneficiaries of this estate and of the '
        'decedent\u2019s surviving spouse, if any, their addresses and relationships to '
        'decedent, and the years of birth of any who are minors, are:',
        keep_with_next=True)
    _beneficiaries_table(doc)

    # 4. Venue
    _pleading_para(doc,
        'Venue of this proceeding is in this county because {venue_reason}.')

    # 5. PR qualifications
    _pleading_para(doc,
        '{pr_name}, whose address is {pr_address}, is qualified to serve as personal '
        'representative of the decedent\u2019s estate: {pr_name} has not been convicted '
        'of a felony, is mentally and physically able to perform the duties of personal '
        'representative, and is 18 years of age or older. '
        '{#pr_is_fl_resident}{pr_name} is a resident of Florida.{/pr_is_fl_resident}'
        '{^pr_is_fl_resident}{pr_name} is not a resident of Florida but is related to '
        'the decedent as {pr_relationship} and is qualified to serve as personal '
        'representative under Florida Statutes section 733.304.{/pr_is_fl_resident}')

    # 6. Petitioner prior conviction
    _pleading_para(doc,
        'Petitioner {#petitioner_has_prior_conviction}has{/petitioner_has_prior_conviction}'
        '{^petitioner_has_prior_conviction}has not{/petitioner_has_prior_conviction} '
        'been convicted in any state or foreign jurisdiction of abuse, neglect, or '
        'exploitation of an elderly person or a disabled adult, as those terms are '
        'defined in Florida Statutes section 825.101.')

    # 7. Preference (was a/b strike alternatives — now a clean conditional)
    _pleading_para(doc,
        '{^higher_preference_exists}No person has equal or higher preference to be '
        'appointed personal representative.{/higher_preference_exists}'
        '{#higher_preference_exists}The following person(s) has/have equal or higher '
        'preference to be appointed personal representative: {higher_preference_names}. '
        '{#higher_preference_formal_notice}Formal notice of this petition will be served '
        'on such person(s).{/higher_preference_formal_notice}'
        '{^higher_preference_formal_notice}Formal notice of this petition will not be '
        'served on such person(s).{/higher_preference_formal_notice}'
        '{/higher_preference_exists}')

    # 8. Assets
    _pleading_para(doc,
        'The nature and approximate value of the assets in this estate are: '
        '{estate_assets_description}.')

    # 9. Estate tax return
    _pleading_para(doc,
        'This estate {#estate_tax_return_required}will{/estate_tax_return_required}'
        '{^estate_tax_return_required}will not{/estate_tax_return_required} be required '
        'to file a federal estate tax return.')

    # 10. Domiciliary proceedings
    _pleading_para(doc,
        'Domiciliary or principal proceedings {#domiciliary_proceedings_pending}are '
        'known to be pending in another state or country. Letters have been issued by '
        '{domiciliary_court_name}, the address of which is {domiciliary_court_address}, '
        'to {domiciliary_representative}, whose address is '
        '{domiciliary_representative_address}.{/domiciliary_proceedings_pending}'
        '{^domiciliary_proceedings_pending}are not known to be pending in another state '
        'or country.{/domiciliary_proceedings_pending}')

    # 11. Will disposition — three-way using will_status field
    _pleading_para(doc,
        '{#will_status_original}The decedent\u2019s last will dated {will_date} {will_year}'
        '{#codicil_dates}, and codicil(s) dated {codicil_dates}{/codicil_dates}, is/are '
        'in the possession of the court or accompanies this petition.{/will_status_original}'
        '{#will_status_authenticated_other}An authenticated copy of a will and/or codicil '
        'deposited with or probated in another jurisdiction accompanies this petition.'
        '{/will_status_authenticated_other}'
        '{#will_status_authenticated_notarial}An authenticated copy of a notarial will or '
        'codicil, the original of which is in the possession of a foreign notary, '
        'accompanies this petition.{/will_status_authenticated_notarial}')

    # 12. No other wills
    _pleading_para(doc,
        'Petitioner is unaware of any unrevoked will or codicil of decedent other than '
        'as set forth in paragraph 11.')

    # Closing — unnumbered.
    indent = Inches(0.5)
    _add_para(doc,
        'Petitioner requests that the decedent\u2019s will be admitted to probate and '
        'that {pr_name} be appointed personal representative of the estate of the decedent.',
        first_indent=indent, space_before=12, space_after=12)
    _add_para(doc,
        'Under penalties of perjury, I declare that I have read the foregoing, and the '
        'facts alleged are true, to the best of my knowledge and belief.',
        first_indent=indent, space_after=18)

    _add_broward_ai_certification(doc, 'Petition for Administration')

    _add_probate_signature_block(doc)

    out_path = os.path.join(TEMPLATE_DIR, 'P3-0100.docx')
    doc.save(out_path)
    _inject_numbering_part(out_path)
    print(f'Wrote {out_path}')


if __name__ == '__main__':
    os.makedirs(TEMPLATE_DIR, exist_ok=True)
    build_p3_0100()
